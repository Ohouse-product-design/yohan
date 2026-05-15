"""Build Product Design TTB budget tracking spreadsheet."""
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment, Border, Font, PatternFill, Side
)
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


# ---------- Styles ----------
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(name="Pretendard", bold=True, color="FFFFFF", size=11)

SUBHEADER_FILL = PatternFill("solid", fgColor="F3F4F6")
SUBHEADER_FONT = Font(name="Pretendard", bold=True, color="111827", size=11)

TITLE_FONT = Font(name="Pretendard", bold=True, size=18, color="111827")
SECTION_FONT = Font(name="Pretendard", bold=True, size=13, color="111827")
NORMAL_FONT = Font(name="Pretendard", size=11, color="111827")
MUTED_FONT = Font(name="Pretendard", size=10, color="6B7280")
BIG_FONT = Font(name="Pretendard", bold=True, size=12, color="111827")

KEY_FILL = PatternFill("solid", fgColor="FEF3C7")
GUIDE_FILL = PatternFill("solid", fgColor="E0F2FE")
INPUT_FILL = PatternFill("solid", fgColor="EFF6FF")
CALC_FILL = PatternFill("solid", fgColor="F9FAFB")

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", indent=1)
RIGHT = Alignment(horizontal="right", vertical="center", indent=1)
WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)

KRW = '#,##0"원";-#,##0"원";'  # zero shown blank
PCT = '0.0%'
DATE = 'yyyy-mm-dd'
PEOPLE = '0"명"'
RATIO = '0.00"x"'


def style_header_row(ws, row, start_col, end_col):
    for c in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# ---------- Build workbook ----------
wb = Workbook()

# =========================================================
# Sheet: 설정
# =========================================================
ws_cfg = wb.active
ws_cfg.title = "설정"

ws_cfg["A1"] = "예산 설정"
ws_cfg["A1"].font = TITLE_FONT
ws_cfg.row_dimensions[1].height = 30

cfg_rows = [
    ("총 예산 (Product Design)", 3699320, KRW, True),
    ("예산 시작월", "2026-01-01", DATE, True),
    ("예산 종료월", "2026-12-31", DATE, True),
    ("기준일 (오늘)", "=TODAY()", DATE, False),
    ("총 개월 수", '=DATEDIF(B3,B4,"M")+1', "0", False),
    ("경과 개월 수",
     '=MIN(MAX(DATEDIF(B3,B5,"M")+1,0), B6)', "0", False),
    ("남은 개월 수", "=MAX(B6-B7,1)", "0", False),
    ("기간 경과율", "=B7/B6", PCT, False),
]

for i, (label, value, fmt, editable) in enumerate(cfg_rows, start=2):
    ws_cfg.cell(row=i, column=1, value=label).font = SUBHEADER_FONT
    ws_cfg.cell(row=i, column=1).fill = SUBHEADER_FILL
    ws_cfg.cell(row=i, column=1).alignment = LEFT
    ws_cfg.cell(row=i, column=1).border = BORDER

    c = ws_cfg.cell(row=i, column=2, value=value)
    c.font = NORMAL_FONT
    c.alignment = RIGHT
    c.number_format = fmt
    c.border = BORDER
    if editable:
        c.fill = INPUT_FILL
    else:
        c.fill = CALC_FILL

ws_cfg["D2"] = "💡 파란색 셀(총 예산, 시작월, 종료월)만 수정하세요. 나머지는 자동 계산됩니다."
ws_cfg["D2"].font = MUTED_FONT

set_col_widths(ws_cfg, [22, 18, 2, 60])

# Named ranges
add_name(wb, "총예산", "설정!$B$2")
add_name(wb, "기준일", "설정!$B$5")
add_name(wb, "총개월", "설정!$B$6")
add_name(wb, "경과개월", "설정!$B$7")
add_name(wb, "남은개월", "설정!$B$8")
add_name(wb, "경과율", "설정!$B$9")


# =========================================================
# Sheet: 멤버 (단순화: 이름 + 비고만)
# =========================================================
ws_mem = wb.create_sheet("멤버")

ws_mem["A1"] = "팀원 명단"
ws_mem["A1"].font = TITLE_FONT
ws_mem.row_dimensions[1].height = 30

ws_mem["A2"] = "💡 이름을 추가/수정하면 사용내역의 이름 드롭다운과 대시보드 사람별 표에 자동 반영됩니다."
ws_mem["A2"].font = MUTED_FONT
ws_mem.merge_cells("A2:C2")

mem_headers = ["이름", "비고", "참고"]
for i, h in enumerate(mem_headers, start=1):
    ws_mem.cell(row=4, column=i, value=h)
style_header_row(ws_mem, 4, 1, 3)
ws_mem.row_dimensions[4].height = 26

default_members = [
    ("이요한(Yohan)", "리더"),
    ("정수현(Gongdee)", ""),
    ("문지선(Sun)", ""),
    ("윤소현(Jenna)", ""),
    ("한수민(Stella)", ""),
    ("이선진(Deeer)", ""),
    ("이재구(Jack)", ""),
    ("조수경(Skamie)", ""),
    ("방채영(Dana)", ""),
    ("윤세린(Selah)", ""),
    ("심희정(Lana)", ""),
    ("백신혜(Gina)", ""),
    ("안유진(Jinny)", ""),
]

MEM_START = 5
MEM_ROWS = 30

for i in range(MEM_ROWS):
    row = MEM_START + i
    if i < len(default_members):
        name, note = default_members[i]
    else:
        name, note = "", ""
    c1 = ws_mem.cell(row=row, column=1, value=name)
    c2 = ws_mem.cell(row=row, column=2, value=note)
    c3 = ws_mem.cell(row=row, column=3)
    for c in (c1, c2, c3):
        c.font = NORMAL_FONT
        c.border = BORDER
        c.fill = INPUT_FILL
        c.alignment = LEFT

set_col_widths(ws_mem, [20, 18, 40])

add_name(
    wb, "멤버목록",
    f"멤버!$A${MEM_START}:$A${MEM_START + MEM_ROWS - 1}"
)


# =========================================================
# Sheet: 사용내역
# =========================================================
ws_log = wb.create_sheet("사용내역")

ws_log["A1"] = "사용 내역 입력"
ws_log["A1"].font = TITLE_FONT
ws_log.row_dimensions[1].height = 30

ws_log["A2"] = (
    "💡 한 줄에 한 건씩 입력. '이름(결제자)'에 결제·기록한 사람을 넣고, "
    "함께한 다른 디자이너는 '함께한 디자이너'에 콤마(,)로 구분해 입력하세요. "
    "총 참여 인원과 1인당 금액은 자동 계산됩니다. "
    "대시보드의 사람별 사용액은 '결제자 + 참여자'에게 1인당 금액으로 균등 분담됩니다."
)
ws_log["A2"].font = MUTED_FONT
ws_log["A2"].alignment = WRAP
ws_log.merge_cells("A2:I2")
ws_log.row_dimensions[2].height = 42

log_headers = [
    "날짜", "이름(결제자)", "카테고리", "사용처/내용", "금액(원)",
    "영수증 링크", "함께한 디자이너", "총 참여 인원", "1인당 금액"
]
for i, h in enumerate(log_headers, start=1):
    ws_log.cell(row=4, column=i, value=h)
style_header_row(ws_log, 4, 1, 9)
ws_log.row_dimensions[4].height = 26

LOG_START = 5
LOG_ROWS = 500
LOG_END = LOG_START + LOG_ROWS - 1

# Sample entries showing both solo and group expenses
samples = [
    ("2026-02-14", "정수현(Gongdee)", "도서", "디자인 시스템 책 구매",
     45000, "", ""),
    ("2026-03-08", "문지선(Sun)", "교육", "Figma 컨퍼런스 티켓",
     180000, "", ""),
    ("2026-03-22", "한수민(Stella)", "툴/구독", "Figma 연간 구독",
     220000, "", ""),
    ("2026-04-10", "이요한(Yohan)", "팀빌딩", "팀 회식 (홍대)",
     360000, "",
     "정수현(Gongdee), 문지선(Sun), 한수민(Stella), 이재구(Jack)"),
]

for i in range(LOG_ROWS):
    row = LOG_START + i

    if i < len(samples):
        date, name, cat, desc, amt, link, joined = samples[i]
        ws_log.cell(row=row, column=1, value=date).number_format = DATE
        ws_log.cell(row=row, column=2, value=name)
        ws_log.cell(row=row, column=3, value=cat)
        ws_log.cell(row=row, column=4, value=desc)
        ws_log.cell(row=row, column=5, value=amt).number_format = KRW
        ws_log.cell(row=row, column=6, value=link)
        ws_log.cell(row=row, column=7, value=joined)
    else:
        ws_log.cell(row=row, column=1).number_format = DATE
        ws_log.cell(row=row, column=5).number_format = KRW

    # H: 총 참여 인원 (자동)
    ws_log.cell(
        row=row, column=8,
        value=(
            f'=IF(B{row}="","",'
            f'IF(TRIM(G{row}&"")="",1,'
            f'LEN(G{row})-LEN(SUBSTITUTE(G{row},",",""))+2))'
        )
    ).number_format = PEOPLE

    # I: 1인당 금액 (자동)
    ws_log.cell(
        row=row, column=9,
        value=(
            f'=IF(OR(B{row}="",E{row}="",H{row}=""),"",'
            f'E{row}/H{row})'
        )
    ).number_format = KRW

    for c in range(1, 10):
        cell = ws_log.cell(row=row, column=c)
        cell.font = NORMAL_FONT
        cell.border = BORDER
        cell.fill = CALC_FILL if c in (8, 9) else INPUT_FILL
        if c in (5, 9):
            cell.alignment = RIGHT
        elif c in (1, 8):
            cell.alignment = CENTER
        else:
            cell.alignment = LEFT

set_col_widths(ws_log, [12, 16, 14, 32, 14, 22, 38, 12, 14])
ws_log.freeze_panes = "A5"

# Data validation: 이름 dropdown
dv_name = DataValidation(
    type="list",
    formula1=f"=멤버!$A${MEM_START}:$A${MEM_START + MEM_ROWS - 1}",
    allow_blank=True,
)
dv_name.error = "멤버 시트에 등록된 이름만 선택할 수 있습니다."
dv_name.errorTitle = "이름 확인"
dv_name.prompt = "멤버 시트에 등록된 이름을 선택하세요."
dv_name.promptTitle = "이름 선택"
ws_log.add_data_validation(dv_name)
dv_name.add(f"B{LOG_START}:B{LOG_END}")

# Data validation: 카테고리 dropdown
dv_cat = DataValidation(
    type="list",
    formula1='"도서,교육,툴/구독,세미나/컨퍼런스,팀빌딩,기타"',
    allow_blank=True,
)
ws_log.add_data_validation(dv_cat)
dv_cat.add(f"C{LOG_START}:C{LOG_END}")

# Named ranges
add_name(wb, "사용_날짜", f"사용내역!$A${LOG_START}:$A${LOG_END}")
add_name(wb, "사용_이름", f"사용내역!$B${LOG_START}:$B${LOG_END}")
add_name(wb, "사용_금액", f"사용내역!$E${LOG_START}:$E${LOG_END}")
add_name(wb, "사용_함께", f"사용내역!$G${LOG_START}:$G${LOG_END}")
add_name(wb, "사용_인원", f"사용내역!$H${LOG_START}:$H${LOG_END}")
add_name(wb, "사용_1인당", f"사용내역!$I${LOG_START}:$I${LOG_END}")


# =========================================================
# Sheet: 대시보드 (재설계)
# =========================================================
ws_dash = wb.create_sheet("대시보드", 0)

ws_dash["A1"] = "Product Design TTB 예산 대시보드"
ws_dash["A1"].font = TITLE_FONT
ws_dash.row_dimensions[1].height = 32
ws_dash.merge_cells("A1:F1")

# --- Section 1: 전체 예산 현황 ---
ws_dash["A3"] = "① 전체 예산 현황"
ws_dash["A3"].font = SECTION_FONT
ws_dash.merge_cells("A3:F3")

# Use SUMIFS-style approach via SUMPRODUCT to safely sum numerics
total_spent_formula = '=SUMPRODUCT((사용_금액<>"")*N(사용_금액))'

basic_rows = [
    ("총 예산", "=총예산", KRW, KEY_FILL),
    ("총 사용액", total_spent_formula, KRW, KEY_FILL),
    ("남은 예산", "=B5-B6", KRW, KEY_FILL),
    ("사용률", "=IFERROR(B6/B5,0)", PCT, KEY_FILL),
    ("기간 경과율", "=경과율", PCT, CALC_FILL),
    ("팀 인원 수",
     f'=COUNTA(멤버!$A${MEM_START}:$A${MEM_START + MEM_ROWS - 1})',
     PEOPLE, CALC_FILL),
]

for i, (label, value, fmt, fill) in enumerate(basic_rows):
    row = 5 + i
    ws_dash.cell(row=row, column=1, value=label).font = SUBHEADER_FONT
    ws_dash.cell(row=row, column=1).fill = SUBHEADER_FILL
    ws_dash.cell(row=row, column=1).alignment = LEFT
    ws_dash.cell(row=row, column=1).border = BORDER

    c = ws_dash.cell(row=row, column=2, value=value)
    c.font = BIG_FONT
    c.alignment = RIGHT
    c.number_format = fmt
    c.border = BORDER
    c.fill = fill

# 사용률 색상 (B8)
ws_dash.conditional_formatting.add(
    "B8",
    CellIsRule(operator="greaterThanOrEqual", formula=["0.9"],
               fill=PatternFill("solid", fgColor="FECACA"),
               font=Font(name="Pretendard", bold=True, color="991B1B"))
)
ws_dash.conditional_formatting.add(
    "B8",
    CellIsRule(operator="between", formula=["0.7", "0.9"],
               fill=PatternFill("solid", fgColor="FEF3C7"),
               font=Font(name="Pretendard", bold=True, color="92400E"))
)

# --- Section 2: 인당 가이드라인 ---
ws_dash["A12"] = "② 인당 사용 가이드라인 (남은 기간 기준)"
ws_dash["A12"].font = SECTION_FONT
ws_dash.merge_cells("A12:F12")

ws_dash["A13"] = (
    "💡 고정 할당이 아닙니다. 팀 비용으로 그때그때 사용할 때, "
    "지금까지·앞으로 인당 어느 정도가 적정한지 보여주는 참고값입니다."
)
ws_dash["A13"].font = MUTED_FONT
ws_dash["A13"].alignment = WRAP
ws_dash.merge_cells("A13:F13")
ws_dash.row_dimensions[13].height = 32

guide_rows = [
    ("📍 현 시점 인당 권장 누적 사용액",
     "=IFERROR(총예산*경과율/B10,0)",
     "지금까지 페이스대로 썼다면 인당 약 이 정도까지 쓴 셈"),
    ("📍 남은 기간 인당 사용 가능액",
     "=IFERROR(B7/B10,0)",
     "남은 예산을 인원수로 나눈 값 (참고)"),
    ("📍 인당 월 권장 사용액 (남은 기간)",
     "=IFERROR(B7/B10/남은개월,0)",
     "남은 기간 동안 한 달에 인당 약 이 정도씩 쓰면 적정"),
    ("📍 팀 전체 월 권장 사용액 (남은 기간)",
     "=IFERROR(B7/남은개월,0)",
     "팀 전체로 한 달에 약 이 정도까지 쓰면 적정"),
]

for i, (label, value, hint) in enumerate(guide_rows):
    row = 15 + i
    ws_dash.cell(row=row, column=1, value=label).font = SUBHEADER_FONT
    ws_dash.cell(row=row, column=1).fill = SUBHEADER_FILL
    ws_dash.cell(row=row, column=1).alignment = LEFT
    ws_dash.cell(row=row, column=1).border = BORDER

    c = ws_dash.cell(row=row, column=2, value=value)
    c.font = BIG_FONT
    c.alignment = RIGHT
    c.number_format = KRW
    c.border = BORDER
    c.fill = GUIDE_FILL

    h = ws_dash.cell(row=row, column=3, value=hint)
    h.font = MUTED_FONT
    h.alignment = LEFT
    ws_dash.merge_cells(start_row=row, start_column=3,
                        end_row=row, end_column=6)

# Named ranges to reference these guideline cells
add_name(wb, "인당_권장_누적", "대시보드!$B$15")
add_name(wb, "인당_남은가능", "대시보드!$B$16")
add_name(wb, "인당_월권장", "대시보드!$B$17")


# --- Section 3: 사람별 사용 현황 ---
ws_dash["A21"] = "③ 사람별 사용 현황 (참여 기준 균등 분담)"
ws_dash["A21"].font = SECTION_FONT
ws_dash.merge_cells("A21:F21")

ws_dash["A22"] = (
    "💡 결제자가 아니어도 '함께한 디자이너'에 이름이 들어간 건은 1인당 금액으로 분담됩니다. "
    "🔴 권장 누적 대비 1.5배 초과 = 빨강, 🟠 1.0~1.5배 = 주황, 🟢 0.5~1.0배 = 초록, ⚪ 0.5배 미만 = 회색."
)
ws_dash["A22"].font = MUTED_FONT
ws_dash["A22"].alignment = WRAP
ws_dash.merge_cells("A22:F22")
ws_dash.row_dimensions[22].height = 42

person_headers = [
    "이름", "사용액", "비중", "평균 대비", "권장 대비", "상태"
]
for i, h in enumerate(person_headers, start=1):
    ws_dash.cell(row=24, column=i, value=h)
style_header_row(ws_dash, 24, 1, 6)
ws_dash.row_dimensions[24].height = 28

PERSON_START = 25
PERSON_END = PERSON_START + MEM_ROWS - 1

# Helper formula chunk: SEARCH-based participant matching
# ","&name&"," within ","&SUBSTITUTE(함께,white,"")&","
def attribution_formula(name_cell):
    """Return formula for person's spend = sum of 1인당 over rows
    where they were payer or in participant list. All spaces are
    stripped on both sides to make matching robust."""
    n = f'SUBSTITUTE({name_cell}," ","")'
    j = 'SUBSTITUTE(사용_함께&""," ","")'
    return (
        f'=IFERROR(SUMPRODUCT('
        f'(사용_이름<>"")*'
        f'(({name_cell}=사용_이름)+'
        f'(({name_cell}<>사용_이름)*'
        f'ISNUMBER(SEARCH(","&{n}&",",","&{j}&","))))*'
        f'IFERROR(N(사용_1인당),0)),0)'
    )

for i in range(MEM_ROWS):
    row = PERSON_START + i
    mem_row = MEM_START + i

    # 이름
    ws_dash.cell(
        row=row, column=1,
        value=f'=IF(멤버!A{mem_row}="","",멤버!A{mem_row})'
    )

    # 사용액 (참여 기준)
    ws_dash.cell(
        row=row, column=2,
        value=(
            f'=IF(A{row}="","",'
            + attribution_formula(f"A{row}")[1:]  # strip leading '='
            + ')'
        )
    ).number_format = KRW

    # 비중 (개인 사용액 / 총 사용액)
    ws_dash.cell(
        row=row, column=3,
        value=f'=IFERROR(IF(A{row}="","",B{row}/$B$6),"")'
    ).number_format = PCT

    # 평균 대비 = 개인 사용액 / (총 사용액 / 인원수)
    ws_dash.cell(
        row=row, column=4,
        value=(
            f'=IFERROR(IF(OR(A{row}="",$B$6=0),"",'
            f'B{row}/($B$6/$B$10)),"")'
        )
    ).number_format = RATIO

    # 권장 대비 = 개인 사용액 / 인당 권장 누적
    ws_dash.cell(
        row=row, column=5,
        value=(
            f'=IFERROR(IF(OR(A{row}="",인당_권장_누적=0),"",'
            f'B{row}/인당_권장_누적),"")'
        )
    ).number_format = RATIO

    # 상태 (텍스트 라벨)
    ws_dash.cell(
        row=row, column=6,
        value=(
            f'=IF(A{row}="","",'
            f'IF(NOT(ISNUMBER(E{row})),"-",'
            f'IF(E{row}>1.5,"🔴 과다",'
            f'IF(E{row}>=1,"🟠 주의",'
            f'IF(E{row}>=0.5,"🟢 정상","⚪ 여유")))))'
        )
    )

    for c in range(1, 7):
        cell = ws_dash.cell(row=row, column=c)
        cell.font = NORMAL_FONT
        cell.border = BORDER
        if c == 1:
            cell.alignment = LEFT
        elif c == 6:
            cell.alignment = CENTER
        else:
            cell.alignment = RIGHT

# Conditional formatting on 권장 대비 (column E)
red_fill = PatternFill("solid", fgColor="FECACA")
red_font = Font(name="Pretendard", bold=True, color="991B1B")
orange_fill = PatternFill("solid", fgColor="FED7AA")
orange_font = Font(name="Pretendard", bold=True, color="9A3412")
green_fill = PatternFill("solid", fgColor="DCFCE7")
green_font = Font(name="Pretendard", color="166534")
gray_fill = PatternFill("solid", fgColor="F3F4F6")
gray_font = Font(name="Pretendard", color="6B7280")

ratio_range = f"E{PERSON_START}:E{PERSON_END}"
ws_dash.conditional_formatting.add(
    ratio_range,
    FormulaRule(
        formula=[f'AND(ISNUMBER(E{PERSON_START}),E{PERSON_START}>1.5)'],
        fill=red_fill, font=red_font)
)
ws_dash.conditional_formatting.add(
    ratio_range,
    FormulaRule(
        formula=[f'AND(ISNUMBER(E{PERSON_START}),E{PERSON_START}>=1,'
                 f'E{PERSON_START}<=1.5)'],
        fill=orange_fill, font=orange_font)
)
ws_dash.conditional_formatting.add(
    ratio_range,
    FormulaRule(
        formula=[f'AND(ISNUMBER(E{PERSON_START}),E{PERSON_START}>=0.5,'
                 f'E{PERSON_START}<1)'],
        fill=green_fill, font=green_font)
)
ws_dash.conditional_formatting.add(
    ratio_range,
    FormulaRule(
        formula=[f'AND(ISNUMBER(E{PERSON_START}),E{PERSON_START}<0.5,'
                 f'E{PERSON_START}>0)'],
        fill=gray_fill, font=gray_font)
)

# Whole row highlight when 권장 대비 > 1.5 (strong visual signal)
row_range = f"A{PERSON_START}:F{PERSON_END}"
ws_dash.conditional_formatting.add(
    row_range,
    FormulaRule(
        formula=[
            f'AND(ISNUMBER($E{PERSON_START}),$E{PERSON_START}>1.5)'
        ],
        fill=PatternFill("solid", fgColor="FEE2E2"))
)

# Data bar on 사용액 column (visual proportion)
ws_dash.conditional_formatting.add(
    f"B{PERSON_START}:B{PERSON_END}",
    DataBarRule(start_type="min", end_type="max",
                color="60A5FA", showValue=True)
)

set_col_widths(ws_dash, [22, 18, 12, 12, 12, 14])
ws_dash.freeze_panes = "A24"


# =========================================================
# Sheet: 카테고리별
# =========================================================
ws_cat = wb.create_sheet("카테고리별")
ws_cat["A1"] = "카테고리별 사용 현황"
ws_cat["A1"].font = TITLE_FONT
ws_cat.row_dimensions[1].height = 30
ws_cat.merge_cells("A1:C1")

cat_headers = ["카테고리", "사용액", "비중"]
for i, h in enumerate(cat_headers, start=1):
    ws_cat.cell(row=3, column=i, value=h)
style_header_row(ws_cat, 3, 1, 3)

categories = ["도서", "교육", "툴/구독", "세미나/컨퍼런스", "팀빌딩", "기타"]
for i, cat in enumerate(categories):
    row = 4 + i
    ws_cat.cell(row=row, column=1, value=cat).font = NORMAL_FONT
    ws_cat.cell(
        row=row, column=2,
        value=(
            f'=SUMIFS(사용_금액,'
            f'사용내역!$C${LOG_START}:$C${LOG_END},A{row})'
        )
    ).number_format = KRW
    ws_cat.cell(
        row=row, column=3,
        value=f'=IFERROR(B{row}/총예산,0)'
    ).number_format = PCT
    for c in range(1, 4):
        cell = ws_cat.cell(row=row, column=c)
        cell.border = BORDER
        cell.font = NORMAL_FONT
        cell.alignment = LEFT if c == 1 else RIGHT

total_row = 4 + len(categories)
ws_cat.cell(row=total_row, column=1, value="합계").font = SUBHEADER_FONT
ws_cat.cell(row=total_row, column=1).fill = SUBHEADER_FILL
ws_cat.cell(row=total_row, column=2,
            value=f'=SUM(B4:B{total_row - 1})').number_format = KRW
ws_cat.cell(row=total_row, column=2).font = SUBHEADER_FONT
ws_cat.cell(row=total_row, column=2).fill = SUBHEADER_FILL
ws_cat.cell(row=total_row, column=3,
            value=f'=IFERROR(B{total_row}/총예산,0)').number_format = PCT
ws_cat.cell(row=total_row, column=3).font = SUBHEADER_FONT
ws_cat.cell(row=total_row, column=3).fill = SUBHEADER_FILL
for c in range(1, 4):
    ws_cat.cell(row=total_row, column=c).border = BORDER
    ws_cat.cell(row=total_row, column=c).alignment = (
        LEFT if c == 1 else RIGHT
    )

set_col_widths(ws_cat, [22, 18, 14])


# Save
out = "/home/user/yohan/product_design_ttb_budget.xlsx"
wb.save(out)
print(f"Saved: {out}")
